# -*- coding: utf-8 -*-
"""智慧职教题目导出（作业/测验/考试）

走 ai.icve.com.cn 与本项目课程列表同一套后端：
- 试卷列表: GET /course/exam/record/getExamListByStudent
- 试卷题目: GET /course/exam/paper?id=&groupId=0
- 作答/答案: GET /course/exam/record/getInfo（按 questionId 对齐，顺序可能被打乱）

考试与测验共用 paper 接口，仅列表 categoryId 不同。
输出：tiku/<课程>/*.md；-dev 时额外写 *.json、raw/、tiku/summary.*
"""

import html as html_module
import json
import re
import time
from pathlib import Path
from typing import Any, Dict, List, Optional

from MoocMain.log import Logger
from AIMoocMain.api import AIMoocApi

# categoryId 映射（以平台实测为准：1作业 / 2考试 / 3测验）
CATEGORY_FALLBACK = {
    1: "作业",
    2: "考试",
    3: "测验",
}

# 答案来源展示名
ANSWER_SOURCE_LABEL = {
    "paper": "试卷接口",
    "exam/record/getInfo": "作答记录接口",
    "taskExamProblemRecordList": "作答记录列表",
}

# 章节测验导出模式
QUIZ_MODE_SINGLE = "single"          # 手动选章节测验
QUIZ_MODE_PER_CHAPTER = "per_chapter"  # 每章一个文件
QUIZ_MODE_MERGED = "merged"          # 全部章节合并一个文件

QUIZ_MODE_LABEL = {
    QUIZ_MODE_SINGLE: "单个/多个章节测验",
    QUIZ_MODE_PER_CHAPTER: "所有章节测验（每章一个文件）",
    QUIZ_MODE_MERGED: "所有章节测验（合并为一个文件）",
}

# 导出文档格式
EXPORT_FORMAT_MD = "md"
EXPORT_FORMAT_DOCX = "docx"
EXPORT_FORMAT_BOTH = "both"

EXPORT_FORMAT_LABEL = {
    EXPORT_FORMAT_MD: "Markdown（.md）",
    EXPORT_FORMAT_DOCX: "Word（.docx）",
    EXPORT_FORMAT_BOTH: "Markdown + Word",
}


def parse_index_input(raw: str, max_n: int) -> Optional[List[int]]:
    """解析课程/测验多选输入：all / 1,3 / 1-3

    返回 1-based 编号列表；非法输入返回 None
    """
    text = (raw or "").strip().lower()
    if not text:
        return None
    if text in ("all", "a", "*", "全部"):
        return list(range(1, max_n + 1))

    picked: List[int] = []
    for part in re.split(r"[,，;；\s]+", text):
        if not part:
            continue
        if "-" in part:
            left, right = part.split("-", 1)
            try:
                start, end = int(left), int(right)
            except ValueError:
                return None
            if start > end:
                start, end = end, start
            picked.extend(range(start, end + 1))
        else:
            try:
                picked.append(int(part))
            except ValueError:
                return None

    if not picked:
        return None
    # 去重保序，过滤越界
    seen = set()
    result = []
    for n in picked:
        if 1 <= n <= max_n and n not in seen:
            seen.add(n)
            result.append(n)
    return result or None


def paper_display_name(item: Dict[str, Any], fallback: str = "") -> str:
    return clean_html_text(_first(
        item, "name", "examName", "title", "paperName", "workExamTitle", default=fallback or "未命名"
    ))


def paper_exam_id(item: Dict[str, Any]) -> str:
    val = _first(item, "id", "examId", "testId", "paperId", "recordId", default="")
    return str(val) if val not in (None, "") else ""


def collect_quiz_catalog(
    client: "AIMoocApi",
    courses: List[Dict[str, Any]],
    category_id: int = 3,
) -> List[Dict[str, Any]]:
    """拉取所选课程的试卷列表，供交互选择章节测验"""
    catalog: List[Dict[str, Any]] = []
    for course in courses:
        course_name = course.get("courseName", "未知课程")
        course_info_id = course.get("id") or course.get("courseInfoId")
        course_id = course.get("courseId")
        if not course_info_id or not course_id:
            continue
        try:
            resp = client.get_exam_list_by_student(
                str(course_info_id), str(course_id), category_id
            )
        except Exception:
            continue
        for item in extract_list(resp):
            exam_id = paper_exam_id(item)
            if not exam_id:
                continue
            catalog.append({
                "course": course_name,
                "course_info_id": str(course_info_id),
                "course_id": str(course_id),
                "category_id": category_id,
                "exam_id": exam_id,
                "paper_name": paper_display_name(item, exam_id),
                "item": item,
            })
    return catalog

# 题型映射（兼容字符串/数字）
QUESTION_TYPE_MAP = {
    1: "单选题", "1": "单选题",
    2: "多选题", "2": "多选题",
    3: "判断题", "3": "判断题",
    4: "填空题", "4": "填空题",
    5: "填空题", "5": "填空题",
    6: "问答题", "6": "问答题",
    7: "匹配题", "7": "匹配题",
    8: "阅读理解", "8": "阅读理解",
    9: "完形填空", "9": "完形填空",
}


def clean_html_text(text: Optional[str]) -> str:
    """去掉 HTML 标签和不可见控制字符，保留纯文本"""
    if not text:
        return ""
    text = str(text)
    # 先把 br/p 换成换行，避免句子粘连
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.I)
    text = re.sub(r'</p\s*>', '\n', text, flags=re.I)
    text = re.sub(r'<[^>]+>', '', text)
    text = html_module.unescape(text)
    for ch in ('​', '‌', '‍', '﻿',
               '‎', '‏', '‪', '‫', '‬', '‭', '‮'):
        text = text.replace(ch, '')
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def safe_filename(name: str) -> str:
    """文件名安全化"""
    name = re.sub(r'[\\/:*?"<>|]+', '_', name or "未命名")
    name = re.sub(r'\s+', '_', name).strip("._")
    return name[:80] or "未命名"


def extract_list(payload: Any) -> List[Dict[str, Any]]:
    """从多种响应结构里抽出列表"""
    if payload is None:
        return []
    if isinstance(payload, list):
        return [x for x in payload if isinstance(x, dict)]
    if isinstance(payload, dict):
        for key in ("rows", "list", "data", "records", "result", "results"):
            val = payload.get(key)
            if isinstance(val, list):
                return [x for x in val if isinstance(x, dict)]
            if isinstance(val, dict):
                for k2 in ("rows", "list", "records", "result"):
                    inner = val.get(k2)
                    if isinstance(inner, list):
                        return [x for x in inner if isinstance(x, dict)]
    return []


def unwrap_paper(payload: Any) -> Dict[str, Any]:
    """从 paper 响应中取出试卷主体"""
    if not isinstance(payload, dict):
        return {}
    if "questions" in payload or "mocPaperDto" in payload:
        return payload
    for key in ("data", "results", "result", "paper"):
        val = payload.get(key)
        if isinstance(val, dict):
            if "questions" in val or "mocPaperDto" in val:
                return val
            for k2 in ("data", "paper", "mocPaperDto"):
                inner = val.get(k2)
                if isinstance(inner, dict):
                    return inner
    return payload


def get_questions_from_paper(paper: Dict[str, Any]) -> List[Dict[str, Any]]:
    """兼容多种试卷结构，取出题目列表"""
    if not paper:
        return []
    if isinstance(paper.get("questions"), list):
        return paper["questions"]
    moc = paper.get("mocPaperDto")
    if isinstance(moc, dict):
        obj = moc.get("objectiveQList") or []
        sub = moc.get("subjectiveQList") or []
        merged = []
        if isinstance(obj, list):
            merged.extend(obj)
        if isinstance(sub, list):
            merged.extend(sub)
        return [x for x in merged if isinstance(x, dict)]
    for key in ("questionList", "problemList", "tmList", "data"):
        val = paper.get(key)
        if isinstance(val, list):
            return [x for x in val if isinstance(x, dict)]
    return []


def _first(d: Dict[str, Any], *keys, default=None):
    for k in keys:
        if k in d and d[k] not in (None, ""):
            return d[k]
    return default


def parse_options(raw: Any) -> List[Dict[str, Any]]:
    """解析选项：兼容 dataJson 字符串 / dataArr / optionDtos / options"""
    if raw is None:
        return []
    if isinstance(raw, str):
        raw = raw.strip()
        if not raw:
            return []
        try:
            raw = json.loads(raw)
        except (json.JSONDecodeError, TypeError):
            return []
    if isinstance(raw, dict):
        # 有时包一层
        for key in ("options", "optionDtos", "dataArr", "dataJson", "list"):
            if key in raw:
                return parse_options(raw.get(key))
        raw = [raw]
    if not isinstance(raw, list):
        return []

    options = []
    for idx, opt in enumerate(raw):
        if isinstance(opt, dict):
            content = clean_html_text(_first(
                opt, "Content", "content", "text", "Text", "title", "name", default=""
            ))
            label = _first(opt, "label", "Label", "SortOrder", "sortOrder", "key", default=None)
            sort_order = _first(opt, "SortOrder", "sortOrder", default=None)
            if label is None:
                label = chr(65 + idx)
            else:
                label = str(label).strip().upper()
                # SortOrder 常是 0/1/2 → A/B/C
                if re.fullmatch(r"\d+", label):
                    label = chr(65 + int(label))
            is_answer = bool(_first(
                opt, "IsAnswer", "isAnswer", "answer", "isRight", "isCorrect", default=False
            ))
            options.append({
                "label": label,
                "text": content,
                "is_answer": is_answer,
                "sort_order": sort_order,
            })
        else:
            options.append({
                "label": chr(65 + idx),
                "text": clean_html_text(str(opt)),
                "is_answer": False,
                "sort_order": idx,
            })
    return options


def _map_answer_tokens(tokens: List[str], options: List[Dict[str, Any]]) -> List[str]:
    """把答案 token 映射成选项 label

    平台 getInfo 的 answer 常是 SortOrder（0/1/2/3），需转成 A/B/C/D。
    """
    if not tokens:
        return []
    sort_to_label = {}
    for opt in options or []:
        so = opt.get("sort_order")
        if so not in (None, ""):
            try:
                sort_to_label[str(int(so))] = opt["label"]
            except (TypeError, ValueError):
                sort_to_label[str(so).strip().upper()] = opt["label"]
    mapped = []
    for t in tokens:
        u = str(t).strip().upper()
        if not u:
            continue
        if u in sort_to_label:
            mapped.append(sort_to_label[u])
        elif re.fullmatch(r"\d+", u):
            idx = int(u)
            if 0 <= idx < len(options or []):
                mapped.append(options[idx]["label"])
            else:
                mapped.append(chr(65 + idx) if idx < 26 else u)
        else:
            mapped.append(u)
    return mapped


def _answer_from_options(options: List[Dict[str, Any]]) -> str:
    """从选项 is_answer 汇总正确答案字母"""
    correct = [o["label"] for o in options if o.get("is_answer")]
    return ",".join(correct) if correct else ""


def _is_judgement_type(type_raw: Any, type_label: str = "") -> bool:
    """是否判断题（typeId=3）"""
    if type_label == "判断题":
        return True
    return type_raw is not None and str(type_raw).strip() in ("3", "判断题")


def _judgement_answer_text(raw: Any) -> str:
    """判断题答案规范成「正确/错误」

    平台约定（实测 record.answer）：1=正确，0=错误。
    兼容 true/false、对/错 等写法；未知原样返回。
    """
    if raw is None:
        return ""
    if isinstance(raw, bool):
        return "正确" if raw else "错误"
    if isinstance(raw, (int, float)):
        if int(raw) == 1:
            return "正确"
        if int(raw) == 0:
            return "错误"
        return str(raw)
    text = clean_html_text(str(raw)).strip()
    if not text:
        return ""
    if text in ("正确", "对", "√", "✓", "是"):
        return "正确"
    if text in ("错误", "错", "×", "否"):
        return "错误"
    u = text.upper()
    if u in ("1", "TRUE", "T", "YES", "Y"):
        return "正确"
    if u in ("0", "FALSE", "F", "NO", "N"):
        return "错误"
    return text


def _normalize_answer(raw: Any, options: Optional[List[Dict[str, Any]]], *, judgement: bool) -> str:
    """按题型规范答案：判断题→正确/错误；其余→A,B,C 字母"""
    if judgement:
        return _judgement_answer_text(raw)
    return _answer_letters(raw, options)


def _answer_letters(raw: Any, options: Optional[List[Dict[str, Any]]] = None) -> str:
    """把 answer 字段规范成 A,B,C 形式

    兼容：
    - "A" / "A,C"
    - "0,2"（SortOrder/下标，有 options 时映射成字母）
    - "[\"A\",\"B\"]"
    """
    if raw is None:
        return ""
    if isinstance(raw, (list, tuple, set)):
        parts = [str(x).strip() for x in raw if str(x).strip() != ""]
    else:
        text = clean_html_text(str(raw))
        if not text:
            return ""
        try:
            parsed = json.loads(text)
            if isinstance(parsed, list):
                parts = [str(x).strip() for x in parsed if str(x).strip() != ""]
            else:
                parts = [p.strip() for p in re.split(r"[,，;；\s]+", text) if p.strip()]
        except (json.JSONDecodeError, TypeError):
            parts = [p.strip() for p in re.split(r"[,，;；\s]+", text) if p.strip()]

    letters = _map_answer_tokens(parts, options or [])
    # 去重保序
    seen = set()
    out = []
    for ch in letters:
        if ch not in seen:
            seen.add(ch)
            out.append(ch)
    return ",".join(out)


def _build_answer_index(info_questions: List[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    """把 record/info 里的题目按多键建索引

    实测 paper.questionId == info.questionId（题库 ID，顺序可能被打乱）
    paper.id == info.paperId（试卷题记录 ID）
    """
    index: Dict[str, Dict[str, Any]] = {}
    for q in info_questions:
        if not isinstance(q, dict):
            continue
        for key in ("questionId", "paperId", "problemRecordId", "id"):
            val = q.get(key)
            if val not in (None, ""):
                index[str(val)] = q
    return index


def _lookup_record_question(
    q: Dict[str, Any],
    index: Dict[str, Dict[str, Any]],
) -> Dict[str, Any]:
    """按 question_id / id / 标题 从 record 索引里取题"""
    for key in ("question_id", "id"):
        val = q.get(key)
        if val not in (None, "") and str(val) in index:
            return index[str(val)]
    # 标题兜底
    title = (q.get("title") or "").strip()
    if title:
        for rec in index.values():
            rec_title = clean_html_text(_first(rec, "title", "content", default=""))
            if rec_title and rec_title == title:
                return rec
    return {}


def _pick_task_ids_for_info(
    paper_record: Dict[str, Any],
    student_records: Optional[List[Dict[str, Any]]] = None,
) -> List[str]:
    """选出 getInfo 可用的 taskId 列表（已交卷优先）

    paper.taskExamRecord 常为 status=0 的未提交草稿，stuAnswer 为空；
    平台「我的答案」在 getStudentRecordByExamId 的已交卷记录里。
    """
    draft_id = str(_first(paper_record, "id", "taskId", "recordId", default="") or "")
    submitted_ids: List[str] = []
    if isinstance(student_records, list):
        submitted = [
            r for r in student_records
            if isinstance(r, dict) and str(r.get("status") or "") == "2"
        ]
        # 多次作答：按提交时间倒序，取最近一次优先
        def _sort_key(r: Dict[str, Any]):
            return str(
                r.get("studentSubmitTime")
                or r.get("updateTime")
                or r.get("createTime")
                or ""
            )

        submitted.sort(key=_sort_key, reverse=True)
        for r in submitted:
            rid = str(r.get("id") or r.get("taskId") or "")
            if rid and rid not in submitted_ids:
                submitted_ids.append(rid)

    task_ids: List[str] = []
    for rid in submitted_ids + ([draft_id] if draft_id else []):
        if rid and rid not in task_ids:
            task_ids.append(rid)
    return task_ids


def _merge_record_answers(
    questions: List[Dict[str, Any]],
    info_questions: List[Dict[str, Any]],
) -> int:
    """把 getInfo/taskExamProblemRecordList 的答案合并进已解析题目

    返回填上 std_answer 或 my_answer 的题数
    """
    if not info_questions:
        return 0
    index = _build_answer_index(info_questions)
    filled = 0

    for q in questions:
        raw = _lookup_record_question(q, index)
        if not raw:
            continue
        options = q.get("options") or []
        judgement = _is_judgement_type(q.get("type_raw"), q.get("type") or "")

        # 正确答案：显式字段优先；getInfo 的 answer 通常就是标准答案
        std = clean_html_text(_first(
            raw,
            "stdAnswer", "correctAnswer", "rightAnswer", "answerContent",
            "trueAnswer", "da", "xsda", "std_answer", "right", default="",
        ))
        if std:
            std = _normalize_answer(std, options, judgement=judgement)

        if not std:
            ans_field = _first(raw, "answer", "trueAnswer", "realAnswer", default=None)
            if ans_field not in (None, ""):
                std = _normalize_answer(ans_field, options, judgement=judgement)

        if not std:
            rec_opts = parse_options(_first(
                raw, "dataJson", "dataArr", "optionDtos", "options", default=None
            ))
            std = _answer_from_options(rec_opts) or _answer_from_options(options)

        # 我的答案：stuAnswer / recordAnswer / myAnswer
        my = clean_html_text(_first(
            raw,
            "stuAnswer", "recordAnswer", "myAnswer", "studentAnswer",
            "userAnswer", "my_answer", "fillBlankRecordAnswer", default="",
        ))
        if my:
            my = _normalize_answer(my, options, judgement=judgement)

        changed = False
        if std and not q.get("std_answer"):
            q["std_answer"] = std
            changed = True
        elif std and q.get("std_answer") and std != q["std_answer"]:
            if len(std) >= len(str(q["std_answer"])):
                q["std_answer"] = std
                changed = True
        if my and not q.get("my_answer"):
            q["my_answer"] = my
            changed = True
        if not q.get("analyse"):
            analyse = clean_html_text(_first(raw, "analysis", "analyse", "analysisText", default=""))
            if analyse:
                q["analyse"] = analyse
                changed = True
        if changed:
            filled += 1

        # 同步选项高亮
        if q.get("std_answer") and q.get("options"):
            labels = {x.strip().upper() for x in str(q["std_answer"]).split(",") if x.strip()}
            for opt in q["options"]:
                if opt.get("label") in labels:
                    opt["is_answer"] = True
    return filled


def parse_question(raw: Dict[str, Any], index: int) -> Dict[str, Any]:
    """把原始题目字段规范成统一结构"""
    type_raw = _first(raw, "typeId", "type", "questionType", "txdm", "qType", default=None)
    type_label = QUESTION_TYPE_MAP.get(type_raw) or QUESTION_TYPE_MAP.get(str(type_raw)) if type_raw is not None else None
    if not type_label:
        type_label = f"题型{type_raw}" if type_raw is not None else "未知"

    title = clean_html_text(_first(
        raw, "title", "questionTitle", "questionContent", "content", "tm", "Title", default=""
    ))

    options = parse_options(_first(
        raw, "dataJson", "dataArr", "optionDtos", "options", "answerList", "xxList", default=None
    ))

    std_answer = clean_html_text(_first(
        raw, "stdAnswer", "correctAnswer", "rightAnswer", "answerContent",
        "trueAnswer", "da", "xsda", default=""
    ))
    judgement = _is_judgement_type(type_raw, type_label)
    if std_answer:
        std_answer = _normalize_answer(std_answer, options, judgement=judgement)

    # paper 未提交时 answer 多为 null；已提交时可能是「我的答案」，不能当标准答案
    my_answer = clean_html_text(_first(
        raw, "stuAnswer", "recordAnswer", "myAnswer", "studentAnswer", "userAnswer", "my_answer", default=""
    ))
    if not my_answer:
        rec = _first(raw, "recordAnswer", "myAnswer", "studentAnswer", "answer", default=None)
        if rec not in (None, ""):
            # 判断题 paper.answer 在未提交时可能为 null，不在此误用 0/1；有 stuAnswer 才填
            if judgement and rec in (0, 1, "0", "1", True, False):
                my_answer = ""
            else:
                my_answer = _normalize_answer(rec, options, judgement=judgement)

    # 选项里已标记正确答案时，汇总到 std_answer
    if not std_answer and options:
        std_answer = _answer_from_options(options)

    question_id = _first(raw, "id", "questionId", "quesId", "tmid", "problemId", default=None)
    score = _first(raw, "score", "questionScore", "scoreValue", "fullScore", default=None)

    return {
        "index": index,
        "id": question_id,
        "question_id": _first(raw, "questionId", default=None),
        "paper_id": _first(raw, "paperId", "problemRecordId", default=None),
        "type_raw": type_raw,
        "type": type_label,
        "title": title,
        "score": score,
        "options": options,
        "std_answer": std_answer,
        "my_answer": my_answer,
        "analyse": clean_html_text(_first(raw, "analyse", "analysis", "explain", "parse", "analysisText", default="")),
    }


def question_to_markdown(q: Dict[str, Any], include_my_answer: bool = False) -> str:
    """单题转 Markdown；include_my_answer 为 True 时才写入我的答案"""
    lines = [f"### {q['index']}. {q['type']}"]
    if q.get("score") not in (None, ""):
        lines[0] += f"（{q['score']}分）"
    lines.append("")
    lines.append(f"**题目：** {q['title'] or '（空）'}")
    lines.append("")
    if q["options"]:
        lines.append("**选项：**")
        lines.append("")
        for opt in q["options"]:
            text = opt.get("text") or ""
            if opt.get("is_answer"):
                # 正确选项整行加粗（不再用 ✅）
                lines.append(f"- **{opt['label']}. {text}**")
            else:
                lines.append(f"- {opt['label']}. {text}")
        lines.append("")
    if q.get("std_answer"):
        lines.append(f"**正确答案：** {q['std_answer']}")
        lines.append("")
    if not q.get("std_answer"):
        lines.append("**答案：** （平台未返回）")
        lines.append("")
    if include_my_answer and q.get("my_answer"):
        lines.append(f"**我的答案：** {q['my_answer']}")
        lines.append("")
    if q.get("analyse"):
        lines.append(f"**解析：** {q['analyse']}")
        lines.append("")
    return "\n".join(lines)


def _question_heading(q: Dict[str, Any]) -> str:
    head = f"{q.get('index', '')}. {q.get('type', '')}"
    if q.get("score") not in (None, ""):
        head += f"（{q['score']}分）"
    return head


# Word 排版：封面页字体较大；正文（含章节）统一小字号
_DOCX_FONT = "微软雅黑"
_DOCX_COVER_TITLE_PT = 16
_DOCX_COVER_META_PT = 12
_DOCX_BODY_PT = 9
_DOCX_CHAPTER_PT = 12


def _apply_docx_font(run, size_pt: float, bold: bool = False) -> None:
    from docx.shared import Pt
    from docx.oxml.ns import qn

    run.font.name = _DOCX_FONT
    run.font.size = Pt(size_pt)
    run.bold = bold
    try:
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), _DOCX_FONT)
        rfonts.set(qn("w:ascii"), _DOCX_FONT)
        rfonts.set(qn("w:hAnsi"), _DOCX_FONT)
    except Exception:
        pass


def _set_normal_style(doc) -> None:
    """Normal 样式：小字号微软雅黑；段间距/行距压紧，提高每页题量。"""
    from docx.shared import Pt

    style = doc.styles["Normal"]
    style.font.name = _DOCX_FONT
    style.font.size = Pt(_DOCX_BODY_PT)
    # Word 默认段后约 8pt、行距 1.15，会把每题撑散
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    try:
        from docx.oxml.ns import qn

        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), _DOCX_FONT)
        rfonts.set(qn("w:ascii"), _DOCX_FONT)
        rfonts.set(qn("w:hAnsi"), _DOCX_FONT)
    except Exception:
        pass


def _configure_section(section) -> None:
    from docx.shared import Cm

    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)


def _add_body_para(
    doc,
    text: str,
    *,
    bold: bool = False,
    keep_with_next: bool = False,
    space_before_pt: float = 0,
):
    """正文段落：小字号；keep_with_next + keep_together，避免题目被分页切割。"""
    from docx.shared import Pt

    p = doc.add_paragraph()
    run = p.add_run(text)
    _apply_docx_font(run, _DOCX_BODY_PT, bold=bold)
    pf = p.paragraph_format
    pf.keep_with_next = keep_with_next
    pf.keep_together = True
    pf.space_before = Pt(space_before_pt)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    return p


def _append_question_to_doc(doc, q: Dict[str, Any], include_my_answer: bool = False) -> None:
    """把一道题写成一组紧凑段落；除最后一段外 keep_with_next，整题尽量同页。

    include_my_answer 为 True 时才写入我的答案。
    """
    # 题间仅留少量空隙；题内段落间距为 0
    gap = 4.0
    parts: List[tuple] = [(_question_heading(q), True, gap)]
    parts.append((q.get("title") or "（空）", True, 0.0))
    if q.get("options"):
        for opt in q["options"]:
            text = opt.get("text") or ""
            parts.append((f"{opt.get('label')}. {text}", bool(opt.get("is_answer")), 0.0))
    my_suffix = f"　|　我的答案：{q['my_answer']}" if include_my_answer and q.get("my_answer") else ""
    if q.get("std_answer"):
        ans = f"正确答案：{q['std_answer']}" + my_suffix
        parts.append((ans, True, 0.0))
    else:
        ans = "答案：（平台未返回）" + my_suffix
        parts.append((ans, False, 0.0))
    if q.get("analyse"):
        parts.append((f"解析：{q['analyse']}", False, 0.0))

    last = len(parts) - 1
    for i, (text, bold, space_before) in enumerate(parts):
        _add_body_para(
            doc,
            text,
            bold=bold,
            keep_with_next=i < last,
            space_before_pt=space_before,
        )


def _set_page_number_start(section, start: int = 1) -> None:
    """正文分节页码从 start 起算（物理第 3 页显示为 1）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    sect_pr = section._sectPr
    for el in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(el)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), str(start))
    sect_pr.append(pg_num)


def _continue_page_number(section) -> None:
    """后续章节分节：去掉复制来的 pgNumType，页码沿用上一节连续编号。"""
    from docx.oxml.ns import qn

    sect_pr = section._sectPr
    for el in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(el)


def _clear_footer(section) -> None:
    """封面/目录页脚不显示页码。"""
    section.footer.is_linked_to_previous = False
    footer = section.footer
    for p in footer.paragraphs:
        for r in list(p.runs):
            r._element.getparent().remove(r._element)


def _add_footer_page_number(section) -> None:
    """页脚居中插入 PAGE 域。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    section.footer.is_linked_to_previous = False
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    _apply_docx_font(run, _DOCX_BODY_PT)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def _set_outline_level(paragraph, level: int = 0) -> None:
    """章节标题写入大纲级别，供 Word 目录域抓取。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    for el in p_pr.findall(qn("w:outlineLvl")):
        p_pr.remove(el)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    p_pr.append(outline)


def _enable_update_fields(doc) -> None:
    """打开文档时自动更新域（目录页码）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    settings = doc.settings.element
    for el in settings.findall(qn("w:updateFields")):
        settings.remove(el)
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    settings.append(el)


def _append_toc(doc, title: str, question_groups: List[Any]) -> None:
    """第 2 页：Word 目录域（章节 + 页码，打开时自动更新）。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    toc_title = doc.add_paragraph()
    run = toc_title.add_run("目录")
    _apply_docx_font(run, _DOCX_COVER_TITLE_PT, bold=True)
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.space_after = Pt(12)

    entries = [h for h, _ in question_groups if h]
    if not entries:
        p = doc.add_paragraph()
        _apply_docx_font(p.add_run(title), _DOCX_COVER_META_PT)
        return

    # TOC 域：\o 标题级别 + \u 大纲级别；dirty 让 Word/WPS 打开时刷新页码
    # 域结果用多段落 + w:br，未更新前也能看到章节列表
    p = doc.add_paragraph()
    r_begin = p.add_run()
    _apply_docx_font(r_begin, _DOCX_COVER_META_PT)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    fld_begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-1" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    r_begin._r.append(fld_begin)
    r_begin._r.append(instr)
    r_begin._r.append(fld_sep)

    r_result = p.add_run()
    _apply_docx_font(r_result, _DOCX_COVER_META_PT)
    for i, heading in enumerate(entries):
        if i > 0:
            br = OxmlElement("w:br")
            r_result._r.append(br)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = heading
        r_result._r.append(t)

    r_end = p.add_run()
    _apply_docx_font(r_end, _DOCX_COVER_META_PT)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r_end._r.append(fld_end)


def write_questions_docx(
    path: Path,
    title: str,
    meta_lines: List[str],
    question_groups: List[Any],
    include_my_answer: bool = False,
) -> None:
    """题目导出为 Word（.docx）

    排版：
    - 第 1 页：封面（标题 + meta），字号较大，无页码
    - 第 2 页：目录（Word TOC 域，章节 + 页码；打开文档时自动更新）
    - 第 3 页起：正文，页码从 1 开始；章节奇数页起、正文字号小
    - 题目各段 keep_with_next / keep_together，空间不够时整题移到下一页

    question_groups: [(section_heading|None, questions), ...]
    """
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION_START
    except ImportError as e:
        raise RuntimeError(
            "缺少 python-docx，无法导出 Word。请执行: uv add python-docx"
        ) from e

    doc = Document()
    _set_normal_style(doc)
    _configure_section(doc.sections[0])
    _clear_footer(doc.sections[0])

    # —— 第 1 页：封面 ——
    cover_title = doc.add_paragraph()
    title_run = cover_title.add_run(title)
    _apply_docx_font(title_run, _DOCX_COVER_TITLE_PT, bold=True)
    cover_title.alignment = 1  # center
    for line in meta_lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        _apply_docx_font(run, _DOCX_COVER_META_PT)

    # —— 第 2 页：目录 ——
    toc_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    _configure_section(toc_section)
    _clear_footer(toc_section)
    _append_toc(doc, title, question_groups)

    # —— 第 3 页起：正文，页码从 1 开始；章节奇数页 ——
    for idx, (section_heading, questions) in enumerate(question_groups):
        section = doc.add_section(WD_SECTION_START.ODD_PAGE)
        _configure_section(section)
        if idx == 0:
            # 物理第 3 页（封面 1 + 目录 2）显示页码 1
            _set_page_number_start(section, 1)
            _add_footer_page_number(section)
        else:
            # 后续章节页码连续，不重置
            _continue_page_number(section)

        if section_heading:
            from docx.shared import Pt

            hp = doc.add_paragraph()
            run = hp.add_run(section_heading)
            _apply_docx_font(run, _DOCX_CHAPTER_PT, bold=True)
            hpf = hp.paragraph_format
            hpf.keep_with_next = True
            hpf.keep_together = True
            hpf.space_before = Pt(0)
            hpf.space_after = Pt(2)
            hpf.line_spacing = 1.0
            # 大纲级别 1：供目录域抓取章节页码
            _set_outline_level(hp, 0)

        for q in questions:
            _append_question_to_doc(doc, q, include_my_answer=include_my_answer)

    _enable_update_fields(doc)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


class QuizExportHandler:
    """按所选课程导出作业/测验/考试题目"""

    def __init__(
        self,
        token: str = None,
        username: str = None,
        password: str = None,
        category_ids: Optional[List[int]] = None,
        output_dir: str = "tiku",
        client: Optional[AIMoocApi] = None,
        course_rows: Optional[List[Dict[str, Any]]] = None,
        course_ids: Optional[List[str]] = None,
        quiz_mode: str = QUIZ_MODE_PER_CHAPTER,
        selected_exam_ids: Optional[List[str]] = None,
        dev_mode: bool = False,
        export_format: str = EXPORT_FORMAT_MD,
        include_my_answer: bool = False,
    ):
        self.logging = Logger(__name__).get_log()
        self.client = client or AIMoocApi(token=token, username=username, password=password)
        self.category_ids = category_ids or [1, 2, 3]
        self.quiz_mode = quiz_mode if 3 in self.category_ids else QUIZ_MODE_PER_CHAPTER
        self.selected_exam_ids = set(selected_exam_ids or [])
        # -dev：额外输出 summary / *.json / raw/；正式模式只写复习用文档
        self.dev_mode = bool(dev_mode)
        # 「我的答案」默认不写入文档，需用户确认后才写
        self.include_my_answer = bool(include_my_answer)
        if export_format not in (EXPORT_FORMAT_MD, EXPORT_FORMAT_DOCX, EXPORT_FORMAT_BOTH):
            export_format = EXPORT_FORMAT_MD
        self.export_format = export_format
        # 题库目录：tiku/<课程名>/，重复导出覆盖同名文件
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        self.stats = {
            "courses": 0,
            "papers": 0,
            "questions": 0,
            "answered": 0,
            "my_answered": 0,
            "errors": 0,
        }
        self.summary_rows: List[Dict[str, Any]] = []

        if course_rows is None:
            courses = self.client.my_course_list(page_size=9999)
            course_rows = courses.get("rows", []) if isinstance(courses, dict) else []
        course_rows = course_rows or []
        if course_ids:
            wanted = {str(x) for x in course_ids}
            course_rows = [
                c for c in course_rows
                if str(c.get("id") or c.get("courseInfoId") or "") in wanted
            ]
        self.course_rows = course_rows
        self.start_export()

    # -------------------- 日志（过程日志仅 -dev） --------------------

    def _info(self, msg: str) -> None:
        """过程/进度日志：仅 -dev 输出；正式模式保持安静"""
        if self.dev_mode:
            self.logging.info(msg)

    def _wants(self, fmt: str) -> bool:
        if self.export_format == EXPORT_FORMAT_BOTH:
            return True
        return self.export_format == fmt

    # -------------------- 主流程 --------------------

    def start_export(self) -> None:
        try:
            self._info(f"📚 待导出课程 {len(self.course_rows)} 门")
            if not self.course_rows:
                self.logging.warning("⚠️ 未选择课程或没有课程数据")
                return

            self._info("➕➕➕ 开始导出题目 ➕➕➕")
            self._info(
                f"⚙️ 测验模式: {QUIZ_MODE_LABEL.get(self.quiz_mode, self.quiz_mode)}"
            )
            self._info(
                f"📄 导出格式: {EXPORT_FORMAT_LABEL.get(self.export_format, self.export_format)}"
            )
            if self.selected_exam_ids:
                self._info(f"🎯 已选章节测验: {len(self.selected_exam_ids)} 份")
            if self.dev_mode:
                self._info("ℹ️ -dev：将额外输出 *.json / raw/ / summary（含姓名/学号等信息，请勿公开分享）")

            for index, course in enumerate(self.course_rows, start=1):
                course_name = course.get("courseName", "未知课程")
                self._info(
                    f"🪧 [{index}] 课程: {course_name} ({course.get('courseInfoName', '')})"
                )
                self.export_course(course)

            self.save_summary()
            self._info("➕➕➕ 导出完成 ➕➕➕")
            self._info(
                f"📊 统计: 课程 {self.stats['courses']} | "
                f"试卷 {self.stats['papers']} | "
                f"题目 {self.stats['questions']} | "
                f"正确答案 {self.stats['answered']} | "
                f"我的答案 {self.stats['my_answered']} | "
                f"错误 {self.stats['errors']}"
            )
            self._info(f"📁 输出目录: {self.output_dir.resolve()}")

        except Exception as e:
            self.logging.error(f"❌ 导出失败: {e}")

    def export_course(self, course: Dict[str, Any]) -> None:
        course_name = course.get("courseName", "未知课程")
        course_info_id = course.get("id") or course.get("courseInfoId")
        course_id = course.get("courseId")

        if not course_info_id or not course_id:
            self.logging.warning(f"⚠️ 课程 {course_name} 缺少 ID，跳过")
            return

        course_dir = self.output_dir / safe_filename(course_name)
        course_dir.mkdir(parents=True, exist_ok=True)

        exported_any = False
        for category_id in self.category_ids:
            try:
                resp = self.client.get_exam_list_by_student(
                    course_info_id, course_id, category_id
                )
            except Exception as e:
                self.logging.error(f"❌ 获取试卷列表失败 course={course_name} cat={category_id}: {e}")
                self.stats["errors"] += 1
                continue

            items = extract_list(resp)
            if not items:
                continue

            category_name = CATEGORY_FALLBACK.get(category_id, str(category_id))
            self._info(
                f"  📋 分类 {category_id}（{category_name}）: {len(items)} 份"
            )

            # 章节测验：按模式过滤/合并
            if category_id == 3 and 3 in self.category_ids:
                if self.quiz_mode == QUIZ_MODE_SINGLE:
                    if not self.selected_exam_ids:
                        self._info(
                            "    ⏭️ 单选模式未选中章节测验，跳过该课程测验"
                        )
                        continue
                    filtered = []
                    for item in items:
                        eid = paper_exam_id(item)
                        if eid and eid in self.selected_exam_ids:
                            filtered.append(item)
                    items = filtered
                    self._info(
                        f"    🎯 按所选章节测验过滤后: {len(items)} 份"
                    )
                if not items:
                    continue

                if self.quiz_mode == QUIZ_MODE_MERGED:
                    if self.export_merged_quizzes(
                        items, category_name, course_name, course_dir,
                        course_info_id=str(course_info_id),
                    ):
                        exported_any = True
                    continue

            for item in items:
                if self.export_one_paper(
                    item, category_id, course_name, course_dir,
                    course_info_id=str(course_info_id),
                ):
                    exported_any = True

        if exported_any:
            self.stats["courses"] += 1

    def load_paper_data(
        self,
        item: Dict[str, Any],
        category_id: int,
        course_name: str,
        course_dir: Path,
        course_info_id: str = "",
    ) -> Optional[Dict[str, Any]]:
        """拉取并解析一份试卷（含答案），不负责写文件"""
        exam_id = paper_exam_id(item)
        if not exam_id:
            return None

        paper_name = paper_display_name(item, exam_id)
        category_name = clean_html_text(_first(
            item, "categoryName", "typeName", "workExamType", default=""
        )) or CATEGORY_FALLBACK.get(category_id, str(category_id))

        try:
            paper_raw = self.client.get_exam_paper(str(exam_id))
        except Exception as e:
            self.logging.error(f"❌ 获取试卷失败 [{category_name}] {paper_name}: {e}")
            self.stats["errors"] += 1
            return None

        paper = unwrap_paper(paper_raw)
        raw_questions = get_questions_from_paper(paper)
        if not raw_questions:
            self.logging.warning(
                f"  ⚠️ 未解析到题目 [{category_name}] {paper_name}（examId={exam_id}）"
            )
            self.save_raw(course_dir, category_name, paper_name, exam_id, paper_raw)
            self.stats["errors"] += 1
            return None

        questions = [parse_question(q, i + 1) for i, q in enumerate(raw_questions)]

        record_info_raw = None
        answer_source = "paper"
        got_answers_from_info = False
        record = paper.get("taskExamRecord") if isinstance(paper, dict) else None
        record = record if isinstance(record, dict) else {}
        draft_task_id = _first(record, "id", "taskId", "recordId", default="")
        rec_course_info_id = _first(
            record, "courseInfoId", default=""
        ) or _first(item, "courseInfoId", default="") or course_info_id
        student_id = (
            _first(record, "userId", "studentId", default="")
            or _first(item, "studentId", default="")
            or getattr(self.client, "user_id", None)
            or ""
        )

        student_records: List[Dict[str, Any]] = []
        if exam_id and student_id:
            try:
                rec_list_raw = self.client.get_exam_student_records(
                    str(exam_id), str(student_id)
                )
                if isinstance(rec_list_raw, list):
                    student_records = [x for x in rec_list_raw if isinstance(x, dict)]
                elif isinstance(rec_list_raw, dict):
                    for key in ("rows", "list", "data", "records"):
                        val = rec_list_raw.get(key)
                        if isinstance(val, list):
                            student_records = [x for x in val if isinstance(x, dict)]
                            break
            except Exception as e:
                self.logging.debug(f"获取作答记录列表失败 examId={exam_id}: {e}")

        task_ids = _pick_task_ids_for_info(record, student_records)
        if rec_course_info_id and task_ids:
            for task_id in task_ids:
                try:
                    record_info_raw = self.client.get_exam_record_info(
                        course_info_id=str(rec_course_info_id),
                        task_id=str(task_id),
                        exam_id=str(exam_id),
                        student_id=str(student_id) if student_id else None,
                    )
                    info_questions = self._extract_info_questions(record_info_raw)
                    if not info_questions:
                        record_info_raw = None
                        continue
                    _merge_record_answers(questions, info_questions)
                    if any(q.get("my_answer") for q in questions):
                        answer_source = "exam/record/getInfo"
                        got_answers_from_info = True
                        break
                    if not got_answers_from_info:
                        answer_source = "exam/record/getInfo"
                        got_answers_from_info = True
                except Exception as e:
                    self.logging.debug(
                        f"获取作答记录失败 examId={exam_id} taskId={task_id}: {e}"
                    )
                    record_info_raw = None

        rec_list = record.get("taskExamProblemRecordList") if isinstance(record, dict) else None
        if isinstance(rec_list, list) and rec_list:
            if _merge_record_answers(questions, rec_list) and not got_answers_from_info:
                answer_source = "taskExamProblemRecordList"

        answered = sum(1 for q in questions if q.get("std_answer"))
        my_answered = sum(1 for q in questions if q.get("my_answer"))
        self.stats["papers"] += 1
        self.stats["questions"] += len(questions)
        self.stats["answered"] = self.stats.get("answered", 0) + answered
        self.stats["my_answered"] = self.stats.get("my_answered", 0) + my_answered

        self.save_raw(course_dir, category_name, paper_name, exam_id, paper_raw)
        if record_info_raw is not None:
            self.save_raw(
                course_dir, category_name, paper_name, exam_id,
                record_info_raw, suffix="_record",
            )

        return {
            "course": course_name,
            "category": category_name,
            "category_id": category_id,
            "paper_name": paper_name,
            "exam_id": exam_id,
            "questions": questions,
            "question_count": len(questions),
            "answered_count": answered,
            "my_answered_count": my_answered,
            "answer_source": answer_source,
            "raw": paper_raw,
            "record_raw": record_info_raw,
        }

    def export_one_paper(
        self,
        item: Dict[str, Any],
        category_id: int,
        course_name: str,
        course_dir: Path,
        course_info_id: str = "",
    ) -> bool:
        data = self.load_paper_data(
            item, category_id, course_name, course_dir, course_info_id
        )
        if not data:
            return False
        self.write_single_paper(course_dir, data)
        return True

    def write_single_paper(self, course_dir: Path, data: Dict[str, Any]) -> None:
        category_name = data["category"]
        paper_name = data["paper_name"]
        exam_id = data["exam_id"]
        questions = data["questions"]
        answered = data["answered_count"]
        my_answered = data["my_answered_count"]
        answer_source = data["answer_source"]

        base = f"{safe_filename(category_name)}_{safe_filename(paper_name)}_{exam_id}"
        payload = {
            "course": data["course"],
            "category": category_name,
            "category_id": data["category_id"],
            "paper_name": paper_name,
            "exam_id": exam_id,
            "question_count": data["question_count"],
            "answered_count": answered,
            "my_answered_count": my_answered,
            "answer_source": answer_source,
            "exported_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "questions": questions,
        }
        json_path = None
        if self.dev_mode:
            json_path = course_dir / f"{base}.json"
            json_path.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
            )

        source_label = ANSWER_SOURCE_LABEL.get(answer_source, answer_source)
        title = f"{data['course']} · {paper_name}"
        meta_lines = [
            f"- 分类: {category_name}",
            f"- 题目数: {len(questions)}",
            f"- 正确答案: {answered}/{len(questions)}",
        ]
        if my_answered > 0:
            meta_lines.append(f"- 我的答案: {my_answered}/{len(questions)}")
        if self.dev_mode:
            meta_lines.append(f"- 答案来源: {source_label}")
            meta_lines.append(f"- 导出时间: {payload['exported_at']}")

        written = []
        md_path = None
        docx_path = None
        if self._wants(EXPORT_FORMAT_MD):
            md_lines = [f"# {title}", ""] + meta_lines + ["", "---", ""]
            for q in questions:
                md_lines.append(question_to_markdown(q, include_my_answer=self.include_my_answer))
                md_lines.append("")
            md_path = course_dir / f"{base}.md"
            md_path.write_text("\n".join(md_lines), encoding="utf-8")
            written.append(md_path.name)
        if self._wants(EXPORT_FORMAT_DOCX):
            docx_path = course_dir / f"{base}.docx"
            write_questions_docx(
                docx_path, title, meta_lines, [(None, questions)],
                include_my_answer=self.include_my_answer,
            )
            written.append(docx_path.name)

        self.summary_rows.append({
            "course": data["course"],
            "category": category_name,
            "paper_name": paper_name,
            "exam_id": exam_id,
            "question_count": len(questions),
            "answered_count": answered,
            "my_answered_count": my_answered,
            "answer_source": answer_source,
            "json": str(json_path) if json_path else "",
            "markdown": str(md_path) if md_path else "",
            "docx": str(docx_path) if docx_path else "",
        })
        self._info(
            f"    ✅ [{category_name}] {paper_name}: {len(questions)} 题 "
            f"（答案 {answered}，我的 {my_answered}） → {', '.join(written) or '（无输出）'}"
        )

    def export_merged_quizzes(
        self,
        items: List[Dict[str, Any]],
        category_name: str,
        course_name: str,
        course_dir: Path,
        course_info_id: str = "",
    ) -> bool:
        """把所选/全部章节测验合并为课程下的一个文件"""
        sections: List[Dict[str, Any]] = []
        for item in items:
            data = self.load_paper_data(
                item, 3, course_name, course_dir, course_info_id
            )
            if not data:
                continue
            # 合并模式下 raw 仍按单卷保存，便于排查
            sections.append(data)

        if not sections:
            return False

        merged_questions: List[Dict[str, Any]] = []
        offset = 0
        section_meta = []
        for sec in sections:
            count = 0
            for q in sec["questions"]:
                q = dict(q)
                q["index"] = offset + count + 1
                q["source_paper"] = sec["paper_name"]
                q["source_exam_id"] = sec["exam_id"]
                merged_questions.append(q)
                count += 1
            offset += count
            section_meta.append({
                "paper_name": sec["paper_name"],
                "exam_id": sec["exam_id"],
                "question_count": count,
                "answered_count": sec["answered_count"],
                "my_answered_count": sec["my_answered_count"],
                "answer_source": sec["answer_source"],
            })

        total = len(merged_questions)
        answered = sum(1 for q in merged_questions if q.get("std_answer"))
        my_answered = sum(1 for q in merged_questions if q.get("my_answer"))

        exported_at = time.strftime("%Y-%m-%d %H:%M:%S")
        base = f"{safe_filename(category_name)}_全部章节测验_{safe_filename(course_name)}"
        payload = {
            "course": course_name,
            "category": category_name,
            "category_id": 3,
            "paper_name": "全部章节测验",
            "quiz_mode": QUIZ_MODE_MERGED,
            "section_count": len(sections),
            "question_count": total,
            "answered_count": answered,
            "my_answered_count": my_answered,
            "exported_at": exported_at,
            "sections": section_meta,
            "questions": merged_questions,
        }
        json_path = None
        if self.dev_mode:
            json_path = course_dir / f"{base}.json"
            json_path.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
            )

        title = f"{course_name} · 全部章节测验"
        meta_lines = [
            f"- 分类: {category_name}",
            f"- 章节数: {len(sections)}",
            f"- 题目数: {total}",
            f"- 正确答案: {answered}/{total}",
        ]
        if my_answered > 0:
            meta_lines.append(f"- 我的答案: {my_answered}/{total}")
        if self.dev_mode:
            meta_lines.append(f"- 答案来源: 作答记录接口/试卷接口（见 summary）")
            meta_lines.append(f"- 导出时间: {exported_at}")

        groups = []
        for sec in sections:
            heading = (
                f"{sec['paper_name']}（{sec['question_count']}题，"
                f"答案 {sec['answered_count']}）"
            )
            groups.append((heading, sec["questions"]))

        written = []
        md_path = None
        docx_path = None
        if self._wants(EXPORT_FORMAT_MD):
            md_lines = [f"# {title}", ""] + meta_lines + ["", "---", ""]
            for heading, questions in groups:
                md_lines.append(f"## {heading}")
                md_lines.append("")
                for q in questions:
                    local_q = dict(q)
                    local_q["index"] = q.get("index", 0)
                    md_lines.append(question_to_markdown(
                        local_q, include_my_answer=self.include_my_answer
                    ))
                    md_lines.append("")
                md_lines.append("---")
                md_lines.append("")
            md_path = course_dir / f"{base}.md"
            md_path.write_text("\n".join(md_lines), encoding="utf-8")
            written.append(md_path.name)
        if self._wants(EXPORT_FORMAT_DOCX):
            docx_path = course_dir / f"{base}.docx"
            write_questions_docx(
                docx_path, title, meta_lines, groups,
                include_my_answer=self.include_my_answer,
            )
            written.append(docx_path.name)

        self.summary_rows.append({
            "course": course_name,
            "category": category_name,
            "paper_name": "全部章节测验",
            "exam_id": "",
            "question_count": total,
            "answered_count": answered,
            "my_answered_count": my_answered,
            "answer_source": "merged",
            "json": str(json_path) if json_path else "",
            "markdown": str(md_path) if md_path else "",
            "docx": str(docx_path) if docx_path else "",
        })
        self._info(
            f"    ✅ [{category_name}] 全部章节测验（{len(sections)}份）: "
            f"{total} 题（答案 {answered}） → {', '.join(written) or '（无输出）'}"
        )
        return True

    @staticmethod
    def _extract_info_questions(record_info_raw: Any) -> List[Dict[str, Any]]:
        """从 getInfo 响应里抽出题目列表"""
        if not isinstance(record_info_raw, dict):
            return []
        data = record_info_raw.get("data") if "data" in record_info_raw else record_info_raw
        if isinstance(data, dict):
            for key in ("questions", "questionList", "paperQuestionList", "examQuestionList", "taskExamProblemRecordList"):
                val = data.get(key)
                if isinstance(val, list) and val:
                    return [x for x in val if isinstance(x, dict)]
            for key in ("paper", "record"):
                inner = data.get(key)
                if isinstance(inner, dict):
                    for key2 in ("questions", "questionList", "paperQuestionList"):
                        val = inner.get(key2)
                        if isinstance(val, list) and val:
                            return [x for x in val if isinstance(x, dict)]
        if isinstance(data, list) and data:
            return [x for x in data if isinstance(x, dict)]
        return []

    def save_raw(
        self,
        course_dir: Path,
        category_name: str,
        paper_name: str,
        exam_id: Any,
        paper_raw: Any,
        suffix: str = "",
    ) -> None:
        """保存原始响应，仅 -dev；接口结构变化时方便排查"""
        if not self.dev_mode:
            return
        raw_dir = course_dir / "raw"
        raw_dir.mkdir(parents=True, exist_ok=True)
        path = raw_dir / (
            f"{safe_filename(category_name)}_{safe_filename(paper_name)}_{exam_id}{suffix}.json"
        )
        try:
            path.write_text(
                json.dumps(paper_raw, ensure_ascii=False, indent=2, default=str),
                encoding="utf-8",
            )
        except Exception as e:
            self.logging.debug(f"保存原始响应失败: {e}")

    def save_summary(self) -> None:
        """summary.json / summary.md，仅 -dev"""
        if not self.dev_mode:
            return
        summary = {
            "exported_at": time.strftime("%Y-%m-%d %H:%M:%S"),
            "output_dir": str(self.output_dir),
            "category_ids": self.category_ids,
            "quiz_mode": self.quiz_mode,
            "quiz_mode_label": QUIZ_MODE_LABEL.get(self.quiz_mode, self.quiz_mode),
            "selected_exam_ids": sorted(self.selected_exam_ids),
            "stats": self.stats,
            "papers": self.summary_rows,
        }
        path = self.output_dir / "summary.json"
        path.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
        )

        lines = [
            "# 题目导出汇总",
            "",
            f"- 时间: {summary['exported_at']}",
            f"- 测验模式: {summary['quiz_mode_label']}",
            "",
        ]
        if not self.summary_rows:
            lines.append("（未导出到任何试卷）")
        for row in self.summary_rows:
            source_label = ANSWER_SOURCE_LABEL.get(
                row.get("answer_source", ""), row.get("answer_source", "")
            )
            lines.append(
                f"- **{row['course']}** / {row['category']} / {row['paper_name']}"
                f"（{row['question_count']}题，答案 {row.get('answered_count', 0)}，"
                f"我的 {row.get('my_answered_count', 0)}，{source_label}）"
            )
        (self.output_dir / "summary.md").write_text("\n".join(lines), encoding="utf-8")
